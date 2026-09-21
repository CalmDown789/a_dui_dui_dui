from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"F:\FPGA预选\10h冲刺")
OUTPUT = ROOT / "PYNQ-Z2开发板申请书.docx"


def set_run_font(run, east_asia: str, latin: str = "Times New Roman") -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7B7B7", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1" if value else "0")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, "宋体")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(90, 90, 90)


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(text)
    return paragraph


def add_body(doc: Document, text: str, *, first_line=True, space_after=4):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Cm(0.74) if first_line else Cm(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.add_run(text)
    return paragraph


def add_bullet(doc: Document, lead: str, text: str):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.first_line_indent = Cm(-0.37)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(lead)
    run.bold = True
    paragraph.add_run(text)
    return paragraph


def format_table_text(table, header=True, body_size=9.5) -> None:
    for r_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    set_run_font(run, "宋体")
                    run.font.size = Pt(body_size)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    if header and r_index == 0:
                        run.bold = True
            if header and r_index == 0:
                set_cell_shading(cell, "D9E2F3")
    if header:
        set_repeat_table_header(table.rows[0])
    set_table_borders(table)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title_style = styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title_style.font.size = Pt(18)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(8)
    title_ppr = title_style.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size in (("Heading 1", 12), ("Heading 2", 11)):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(3)

    list_style = styles["List Bullet"]
    list_style.font.name = "Times New Roman"
    list_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    list_style.font.size = Pt(10.2)
    list_style.font.color.rgb = RGBColor(0, 0, 0)
    list_style.paragraph_format.line_spacing = 1.12

    title = doc.add_paragraph(style="Title")
    title.add_run("PYNQ Z2 开发板申请书")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(7)
    run = subtitle.add_run("东南大学第二十届 PLD 设计竞赛")
    set_run_font(run, "宋体")
    run.font.size = Pt(11)
    run.font.bold = True

    info = doc.add_table(rows=3, cols=4)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    widths = [Cm(2.5), Cm(5.2), Cm(2.5), Cm(5.2)]
    for row in info.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
    values = [
        ("团队名称", "________________", "联系人", "________________"),
        ("团队成员", "________________", "联系方式", "________________"),
        ("申请设备", "PYNQ-Z2 开发板 1 块", "使用期限", "领板至竞赛验收结束"),
    ]
    for row, content in zip(info.rows, values):
        for cell, value in zip(row.cells, content):
            cell.text = value
    for row in info.rows:
        for idx in (0, 2):
            set_cell_shading(row.cells[idx], "E7E6E6")
            for run in row.cells[idx].paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(10)
    format_table_text(info, header=False, body_size=10)

    add_heading(doc, "一 申请事项")
    add_body(
        doc,
        "我队拟参加东南大学第二十届 PLD 设计竞赛，自主命题为“基于 PYNQ-Z2 的轻量视频超分 FPGA 加速系统”。现申请 PYNQ-Z2 开发板 1 块，用于完成加速器 IP 集成、PS 与 PL 协同控制、图像数据搬运、板上验证和最终演示。该板卡不是通用学习替代品，而是本项目既定系统架构中的必要组成。",
    )

    add_heading(doc, "二 赛题内容与拟完成工作")
    add_body(
        doc,
        "本项目面向低带宽视频回传、监控与移动端视觉等场景，对 640×360 单通道 Y 图像进行 2 倍超分处理，输出 1280×720 图像。系统采用三层 3×3 卷积网络，通道结构为 1→8→16→4，前两层使用 PReLU，末端通过 Pixel Shuffle 完成 2 倍上采样；权重与激活采用 INT8，卷积累加采用 INT32。",
    )
    add_bullet(doc, "算法与数据：", "完成轻量模型训练或微调、INT8 量化、逐层参数导出、黄金结果生成，并以双三次插值作为对照。")
    add_bullet(doc, "PL 加速器：", "使用 Vitis HLS 实现流式行缓存、3×3 滑窗、并行乘加、跨通道累加、定点后处理和 Pixel Shuffle，完成 C 仿真与综合。")
    add_bullet(doc, "PS 与系统集成：", "由 PS 侧加载 overlay、配置 AXI4-Lite 寄存器、控制 DMA 或视频缓冲区，并完成输入输出、结果读取和显示。")
    add_bullet(doc, "验证与演示：", "完成 Python 黄金模型和硬件输出的逐层或端到端比对，记录资源、时序、处理耗时、PSNR 与 SSIM，优先实现板上单帧闭环并争取连续帧演示。")

    add_heading(doc, "三 必须使用 PYNQ Z2 的原因")
    add_body(
        doc,
        "本项目不是孤立的卷积 RTL 验证，而是“图像输入—PS 调度—PL 加速—结果回读—显示与评价”的完整系统。PYNQ-Z2 提供的 Zynq PS、PYNQ 软件栈和板级视频接口同时参与上述链路，缺少其中任一部分都会显著改写方案。",
    )

    reasons = doc.add_table(rows=1, cols=3)
    reasons.alignment = WD_TABLE_ALIGNMENT.CENTER
    reasons.autofit = False
    header = reasons.rows[0].cells
    header[0].text = "项目必需功能"
    header[1].text = "PYNQ-Z2 的具体作用"
    header[2].text = "替代平台带来的问题"
    reason_rows = [
        (
            "PS 与 PL 协同",
            "双核 Cortex-A9 运行 PYNQ/Linux 和 Python，负责 overlay 加载、寄存器配置、DMA 调度与结果管理。",
            "纯 Artix-7 板无硬核 ARM，需额外实现 MicroBlaze 或全硬件控制，系统架构与开发工作量均发生根本变化。",
        ),
        (
            "视频与显示闭环",
            "板载 HDMI 输入输出及 PYNQ base overlay 可作为视频采集、缓冲和显示基础，使加速器能进入真实数据通路。",
            "ACX750-200T 无 HDMI 输入；其他 Zynq 板的板级接口与 overlay 不等同，现有 PYNQ-Z2 设计不能直接迁移。",
        ),
        (
            "DDR 与高速搬运",
            "PS 侧硬核 DDR 控制器配合 AXI DMA/VDMA 搬运帧数据，Python 只承担控制，避免逐像素软件访问。",
            "纯 PL 平台需另行调通 MIG、软核和等效视频缓冲链路，评估工作量约为 PYNQ 路线的 3 至 5 倍。",
        ),
        (
            "卷积资源余量",
            "XC7Z020 提供 220 个 DSP48E1 与约 630 KB BRAM，可容纳 INT8 卷积并行单元、权重和多通道行缓存。",
            "7015 仅有 160 个 DSP 和约 427 KB BRAM，资源余量更小；本项目还需为 AXI、DMA 与视频链路预留资源。",
        ),
    ]
    for values in reason_rows:
        cells = reasons.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    for row in reasons.rows:
        row.cells[0].width = Cm(2.8)
        row.cells[1].width = Cm(6.0)
        row.cells[2].width = Cm(6.4)
    format_table_text(reasons, header=True, body_size=8.3)
    for row in reasons.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=65, start=95, bottom=65, end=95)

    doc.add_page_break()

    add_heading(doc, "四 已有基础与板卡匹配性")
    add_body(
        doc,
        "团队已按 PYNQ-Z2 对应器件 XC7Z020-1CLG400C 建立 Vitis HLS 工程，而非在收到板卡后才开始准备。目前已完成 3×3 流式窗口、多通道卷积、定点后处理和 Pixel Shuffle 的独立验证，并完成末层“卷积—后处理—Pixel Shuffle”数据流集成。两个整链路 C 仿真用例通过，HLS 已识别出三个并行进程。",
    )
    add_body(
        doc,
        "当前末层子系统的 HLS 估算为 20 个 BRAM_18K、11 个 DSP、3899 个 FF、4972 个 LUT，估算最高频率约 140.87 MHz，并已导出可供 Vivado 检查端口与搭建占位 Block Design 的预览 IP。上述数据仅用于说明设计已针对 PYNQ-Z2 开展并具有初步可综合性，最终资源、时序和帧率将以完整系统实现及板上实测为准。",
    )

    add_heading(doc, "五 领板后的实施计划")
    plan = doc.add_table(rows=1, cols=3)
    plan.alignment = WD_TABLE_ALIGNMENT.CENTER
    plan.autofit = False
    plan.rows[0].cells[0].text = "阶段"
    plan.rows[0].cells[1].text = "主要工作"
    plan.rows[0].cells[2].text = "验收产物"
    plan_rows = [
        ("环境验收", "刷写并核对 PYNQ 镜像，验证 Jupyter、网络、HDMI 与基础 overlay。", "环境记录与基线画面"),
        ("硬件集成", "将 HLS IP 接入同一 overlay，完成 AXI4-Stream、AXI4-Lite、DMA/VDMA 和地址连接。", "bitstream、hwh 与 Block Design"),
        ("位精确联调", "装载真实 INT8 权重和测试向量，对比 Python 黄金模型，定位首个不一致层。", "对拍结果与调试记录"),
        ("系统测评", "运行 640×360 到 1280×720 单帧闭环，测量时序、资源、耗时、PSNR 与 SSIM。", "实测表格、截图与演示视频"),
    ]
    for values in plan_rows:
        cells = plan.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    for row in plan.rows:
        row.cells[0].width = Cm(2.4)
        row.cells[1].width = Cm(8.2)
        row.cells[2].width = Cm(4.6)
    format_table_text(plan, header=True, body_size=9)

    add_heading(doc, "六 预期成果")
    add_bullet(doc, "硬件成果：", "可综合的轻量超分 HLS IP、完整 overlay、约束与综合实现报告。")
    add_bullet(doc, "软件成果：", "INT8 黄金模型、测试向量、PYNQ Python 驱动和自动化验证脚本。")
    add_bullet(doc, "演示成果：", "至少完成板上单帧超分与 Jupyter 对比展示；在接口条件允许时进一步完成 HDMI 或连续帧演示。")
    add_bullet(doc, "评价成果：", "以真实报告和板上测量记录资源、时序、处理耗时、PSNR 与 SSIM，不使用未经验证的性能数据。")

    add_heading(doc, "七 使用承诺")
    add_body(
        doc,
        "我队承诺仅将申请的 PYNQ-Z2 用于本届 PLD 竞赛项目开发与验收，遵守板卡领用、保管和归还要求；领板时携带本申请书纸质版，并按要求记录开发过程、提交可复现工程与演示材料。项目结束后按赛事安排及时归还设备。",
    )

    signature = doc.add_table(rows=3, cols=2)
    signature.alignment = WD_TABLE_ALIGNMENT.RIGHT
    signature.autofit = False
    signature.cell(0, 0).text = "团队负责人签字："
    signature.cell(0, 1).text = "________________"
    signature.cell(1, 0).text = "团队成员签字："
    signature.cell(1, 1).text = "________________"
    signature.cell(2, 0).text = "申请日期："
    signature.cell(2, 1).text = "______年____月____日"
    for row in signature.rows:
        row.cells[0].width = Cm(3.6)
        row.cells[1].width = Cm(5.0)
        for cell in row.cells:
            set_cell_margins(cell, top=65, start=80, bottom=65, end=80)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, "宋体")
                    run.font.size = Pt(10.5)
    # Intentionally borderless signature area.
    tbl_pr = signature._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)

    footer = section.footer
    add_page_number(footer.paragraphs[0])

    doc.core_properties.title = "PYNQ Z2 开发板申请书"
    doc.core_properties.subject = "东南大学第二十届 PLD 设计竞赛开发板申请"
    doc.core_properties.author = "项目团队"
    doc.core_properties.keywords = "PYNQ-Z2, PLD, FPGA, 视频超分"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
